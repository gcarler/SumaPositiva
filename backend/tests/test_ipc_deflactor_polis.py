"""El IPC oficial solo deflacta flujos anuales completos y conserva el anexo."""

from __future__ import annotations

import hashlib
import json
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.modules.compromisos.ipc_deflactor_polis import (
    MESES,
    ORIGEN,
    SHA_ANEXO_SEP2026,
    cargar_factores,
    convertir_flujo_anual,
    extraer_indices,
    factor_desde_diciembre_2019,
    factor_hacia_anio_base,
    materializar_deflactor,
)


def _anexo(ruta: Path) -> Path:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "IndicesIPC"
    hoja.cell(8, 25, "Base Diciembre de 2018 = 100,00")
    hoja.cell(9, 1, "Mes")
    for columna, anio in enumerate(range(2012, 2027), start=2):
        hoja.cell(9, columna, anio)
        for mes, nombre in enumerate(MESES, start=1):
            hoja.cell(9 + mes, 1, nombre)
            if anio != 2026 or mes <= 9:
                hoja.cell(9 + mes, columna, 200 if anio == 2024 else 220 if anio == 2026 else 100)
    libro.save(ruta)
    return ruta


def test_ipc_dane_deflacta_anos_completos_y_excluye_2026(tmp_path: Path) -> None:
    origen = _anexo(tmp_path / "anexo.xlsx")
    salida = materializar_deflactor(origen, tmp_path / "salidas")
    datos = json.loads((tmp_path / "salidas" / f"{salida['sha256']}.json").read_text())
    por_anio = {fila["anio"]: fila for fila in datos["anios"]}
    assert salida["anios_completos"] == 14
    assert Decimal(por_anio[2012]["factor_cop_promedio_2024"]) == 2
    assert Decimal(por_anio[2024]["factor_cop_promedio_2024"]) == 1
    assert por_anio[2012]["factor_mostrado_8d"] == "2.00000000"
    assert por_anio[2026]["factor_cop_promedio_2024"] is None
    assert por_anio[2026]["meses_disponibles"] == 9
    assert convertir_flujo_anual(Decimal("100.00"), 2012, {2012: Decimal(2)}) == Decimal("200.00")
    with pytest.raises(ValueError, match="Sin IPC anual completo"):
        convertir_flujo_anual(Decimal(100), 2026, {2012: Decimal(2)})
    assert materializar_deflactor(origen, tmp_path / "salidas") == salida


def test_ipc_rechaza_anexo_incompleto(tmp_path: Path) -> None:
    origen = _anexo(tmp_path / "anexo.xlsx")
    from openpyxl import load_workbook

    libro = load_workbook(origen)
    libro.active.cell(21, 2, None)
    libro.active.cell(21, 2).value = None
    libro.save(origen)
    with pytest.raises(ValueError, match="2012 está incompleta"):
        extraer_indices(origen)


def test_anexo_dane_real_conserva_base_y_hash() -> None:
    assert hashlib.sha256(ORIGEN.read_bytes()).hexdigest() == SHA_ANEXO_SEP2026
    indices = extraer_indices(ORIGEN)
    assert indices[2018][12] == Decimal("100")
    assert indices[2025][12] == Decimal("152.27")
    assert len(indices[2026]) == 9


def test_factores_reales_verificados_no_incluyen_anio_parcial(tmp_path: Path) -> None:
    resultado = materializar_deflactor(destino=tmp_path / "deflactor")
    ruta = Path(__file__).resolve().parents[2] / resultado["ruta"]
    factores = cargar_factores(ruta)
    assert factores[2024] == 1
    assert 2026 not in factores
    assert convertir_flujo_anual(Decimal("100"), 2012, factores) > Decimal("100")
    assert factor_hacia_anio_base(2025, 2025) == 1
    assert factor_hacia_anio_base(2024, 2025) is not None
    assert factor_hacia_anio_base(2026, 2025) is None
    assert factor_hacia_anio_base(2024, 2027) is None
    assert factor_desde_diciembre_2019(2019) > 0
    assert factor_desde_diciembre_2019(2025) > factor_desde_diciembre_2019(2019)
    assert factor_desde_diciembre_2019(2026) is None
    alterada = tmp_path / ruta.name
    alterada.write_bytes(ruta.read_bytes() + b" ")
    with pytest.raises(ValueError, match="SHA-256"):
        cargar_factores(alterada)
    forjada = json.loads(ruta.read_text())
    forjada["anios"][0]["factor_cop_promedio_2024"] = "999"
    contenido = (json.dumps(forjada, sort_keys=True) + "\n").encode()
    ruta_forjada = tmp_path / f"{hashlib.sha256(contenido).hexdigest()}.json"
    ruta_forjada.write_bytes(contenido)
    with pytest.raises(ValueError, match="no está autorizada"):
        cargar_factores(ruta_forjada)

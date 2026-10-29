"""Deflactor de flujos anuales COP desde la serie empalmada de IPC nacional DANE."""

from __future__ import annotations

import hashlib
import json
from decimal import ROUND_HALF_UP, Decimal
from functools import lru_cache
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.modules.compromisos.series_7d_polis import RAIZ, _escribir_inmutable

ORIGEN = RAIZ / "data/banco_fuentes_oficiales/anex-IPC-Indices-sep2026.xlsx"
MANIFIESTO_ORIGEN = ORIGEN.with_suffix(".manifest.json")
DESTINO = RAIZ / "data/capas_analiticas/ipc_deflactor"
URL_DANE = "https://www.dane.gov.co/files/operaciones/IPC/sep2026/anex-IPC-Indices-sep2026.xlsx"
SHA_ANEXO_SEP2026 = "56f8e5e6e9d235d411e487b2284beba1be8eaeee69b7b9516bfa8ea8443c0676"
SHA_DEFLECTOR_V2 = "49e44c7feee0f72b92f80dd03d388509681a2597a48f502f686349504810dc57"
MESES = (
    "Enero",
    "Febrero",
    "Marzo",
    "Abril",
    "Mayo",
    "Junio",
    "Julio",
    "Agosto",
    "Septiembre",
    "Octubre",
    "Noviembre",
    "Diciembre",
)


def _decimal(valor: Any) -> Decimal:
    if isinstance(valor, bool) or not isinstance(valor, (int, float, Decimal)):
        raise ValueError("El índice IPC contiene un valor no numérico")
    numero = Decimal(str(valor))
    if not numero.is_finite() or numero <= 0:
        raise ValueError("El índice IPC debe ser positivo y finito")
    return numero


def extraer_indices(ruta: Path) -> dict[int, dict[int, Decimal]]:
    """Lee índices mensuales de la hoja nacional; no infiere los meses ausentes."""
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        if libro.sheetnames != ["IndicesIPC"]:
            raise ValueError("El anexo DANE cambió de hojas")
        hoja = libro["IndicesIPC"]
        if "Base Diciembre de 2018 = 100,00" not in str(hoja.cell(8, 25).value):
            raise ValueError("La base del IPC no es diciembre de 2018")
        encabezado = [celda.value for celda in hoja[9]]
        columnas = {int(anio): indice for indice, anio in enumerate(encabezado) if isinstance(anio, int)}
        if not set(range(2012, 2027)).issubset(columnas):
            raise ValueError("Faltan años 2012–2026 en la serie DANE")
        indices: dict[int, dict[int, Decimal]] = {anio: {} for anio in range(2012, 2027)}
        for mes, nombre in enumerate(MESES, start=1):
            fila = [celda.value for celda in hoja[9 + mes]]
            if fila[0] != nombre:
                raise ValueError(f"Orden de meses inesperado: {fila[0]}")
            for anio, valores in indices.items():
                valor = fila[columnas[anio]]
                if valor is not None:
                    valores[mes] = _decimal(valor)
        if indices[2018][12] != Decimal("100"):
            raise ValueError("Diciembre de 2018 no coincide con la base 100")
        for anio in range(2012, 2026):
            if len(indices[anio]) != 12:
                raise ValueError(f"La serie IPC anual {anio} está incompleta")
        if set(indices[2026]) != set(range(1, 10)):
            raise ValueError("El corte de septiembre de 2026 cambió; revisar anexo")
        return indices
    finally:
        libro.close()


def convertir_flujo_anual(monto_corriente: Decimal, anio: int, factores: dict[int, Decimal]) -> Decimal:
    """Convierte un flujo anual a COP promedio de 2024; 2026 parcial se rechaza."""
    if anio not in factores:
        raise ValueError(f"Sin IPC anual completo para {anio}")
    if not monto_corriente.is_finite() or monto_corriente < 0:
        raise ValueError("El monto COP debe ser finito y no negativo")
    return (monto_corriente * factores[anio]).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def cargar_factores(ruta: Path) -> dict[int, Decimal]:
    """Carga únicamente factores de años completos desde un artefacto íntegro."""
    contenido = ruta.read_bytes()
    if ruta.stem != hashlib.sha256(contenido).hexdigest():
        raise ValueError("El nombre del deflactor no coincide con su SHA-256")
    if ruta.stem != SHA_DEFLECTOR_V2:
        raise ValueError("La versión del deflactor no está autorizada")
    datos = json.loads(contenido)
    if (
        datos.get("esquema") != "polis.ipc_deflactor_flujos_anuales.v2"
        or datos.get("base_pesos_constantes") != "promedio_anual_2024"
        or datos.get("fuente_sha256") != SHA_ANEXO_SEP2026
    ):
        raise ValueError("Deflactor no autorizado para COP promedio de 2024")
    factores = {
        int(fila["anio"]): Decimal(fila["factor_cop_promedio_2024"])
        for fila in datos["anios"]
        if fila["estado"] == "anio_completo"
    }
    if set(factores) != set(range(2012, 2026)) or factores[2024] != 1:
        raise ValueError("La serie anual del deflactor está incompleta")
    return factores


@lru_cache(maxsize=1)
def _factores_oficiales() -> dict[int, Decimal]:
    return cargar_factores(DESTINO / f"{SHA_DEFLECTOR_V2}.json")


def factor_hacia_anio_base(anio: int, base: int, ruta: Path | None = None) -> Decimal | None:
    """Factor de un flujo anual completo hacia COP promedio del año base."""
    factores = cargar_factores(ruta) if ruta is not None else _factores_oficiales()
    if anio not in factores or base not in factores:
        return None
    return factores[anio] / factores[base]


@lru_cache(maxsize=1)
def _indices_oficiales() -> dict[int, dict[int, Decimal]]:
    if hashlib.sha256(ORIGEN.read_bytes()).hexdigest() != SHA_ANEXO_SEP2026:
        raise ValueError("El anexo IPC DANE cambió sin nueva versión")
    return extraer_indices(ORIGEN)


def factor_desde_diciembre_2019(anio_destino: int) -> Decimal | None:
    """Pasa COP de diciembre de 2019 a COP promedio anual del destino."""
    indices = _indices_oficiales()
    destino = indices.get(anio_destino, {})
    if len(destino) != 12:
        return None
    return (sum(destino.values()) / Decimal(12)) / indices[2019][12]


def materializar_deflactor(origen: Path = ORIGEN, destino: Path = DESTINO) -> dict[str, Any]:
    """Crea factores reproducibles base 2024 con SHA del anexo oficial."""
    sha_origen = hashlib.sha256(origen.read_bytes()).hexdigest()
    if origen.resolve() == ORIGEN.resolve():
        metadatos = json.loads(MANIFIESTO_ORIGEN.read_text(encoding="utf-8"))
        if (
            sha_origen != SHA_ANEXO_SEP2026
            or metadatos.get("esquema") != "polis.fuente_transversal.v1"
            or metadatos.get("sha256_archivo") != sha_origen
            or metadatos.get("url_archivo") != URL_DANE
            or metadatos.get("fecha_actualizacion_dane") != "2026-10-07"
        ):
            raise ValueError("El anexo oficial o su manifiesto cambió sin versionar origen y fecha")
    indices = extraer_indices(origen)
    promedios = {anio: sum(valores.values()) / Decimal(12) for anio, valores in indices.items() if len(valores) == 12}
    base = promedios[2024]
    factores = {anio: base / promedio for anio, promedio in promedios.items()}
    anios = []
    for anio, valores in sorted(indices.items()):
        completo = len(valores) == 12
        anios.append(
            {
                "anio": anio,
                "meses_disponibles": len(valores),
                "ultimo_mes": MESES[max(valores) - 1],
                "ipc_promedio_anual": str(promedios[anio].quantize(Decimal("0.000001"))) if completo else None,
                "ipc_diciembre": str(valores[12]) if 12 in valores else None,
                "factor_cop_promedio_2024": str(factores[anio]) if completo else None,
                "factor_mostrado_8d": (str(factores[anio].quantize(Decimal("0.00000001"))) if completo else None),
                "estado": "anio_completo" if completo else "parcial_no_deflactar_flujo_anual",
            }
        )
    salida = {
        "esquema": "polis.ipc_deflactor_flujos_anuales.v2",
        "fuente": "DANE, IPC total nacional, índices serie de empalme",
        "url_fuente": URL_DANE,
        "fecha_actualizacion_dane": "2026-10-07",
        "fuente_sha256": sha_origen,
        "base_indice": "diciembre_2018_igual_100",
        "base_pesos_constantes": "promedio_anual_2024",
        "metodo": "IPC promedio de los 12 meses de 2024 dividido por IPC promedio de los 12 meses del año origen",
        "redondeo": "cociente Decimal con 28 dígitos de precisión; vista a 8 decimales y COP a centavos",
        "anios": anios,
        "estado_publicacion": "no_publicable",
        "limitacion": "Solo flujos anuales completos; 2026 y asignaciones multianuales requieren corte mensual.",
    }
    contenido = (json.dumps(salida, ensure_ascii=False, sort_keys=True, indent=2) + "\n").encode()
    sha = hashlib.sha256(contenido).hexdigest()
    if origen.resolve() == ORIGEN.resolve() and sha != SHA_DEFLECTOR_V2:
        raise ValueError("La derivación IPC cambió sin versionar su método")
    destino.mkdir(parents=True, exist_ok=True)
    ruta = destino / f"{sha}.json"
    _escribir_inmutable(ruta, contenido)
    return {
        "ruta": ruta.relative_to(RAIZ).as_posix() if ruta.is_relative_to(RAIZ) else str(ruta),
        "sha256": sha,
        "fuente_sha256": sha_origen,
        "anios_completos": len(promedios),
        "estado_publicacion": "no_publicable",
    }
